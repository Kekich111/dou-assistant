"""Проверка загруженного приказа / распоряжения (без LLM).

Сверяет DOCX с модулем (по методрекомендациям) и Инструкцией, находит незаполненные
поля, лишние служебные элементы, ошибки в датах и наименованиях и формирует
дальнейшие шаги со ссылками на пункты.
"""
from __future__ import annotations

import io
import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from app.rules.fill import DATE_IN_TEXT, MONTHS, _parse_date

BLANK = re.compile(r"_{3,}")
MARK = re.compile(r"\s*\[\d+\]")
VARIANT = re.compile(r"([а-яё]+) / ([а-яё]+)")
QUOTED = re.compile(r"«([^»]{1,80})»")
TEXT_DATE = re.compile(r"«(\d{1,2})»\s+([а-яё]+)\s+(\d{4})")

FREE_PARAGRAPHS = {"НАИМЕНОВАНИЕ ОРГАНИЗАЦИИ", "г. Москва"}   # подставляются свои
NOTES_HEAD = "Примечания по заполнению"

MR = "Методические рекомендации"
INS = "Инструкция"


def _n(s: str) -> str:
    return re.sub(r"\s+", " ", s.replace("\xa0", " ")).strip()


def _clean(s: str) -> str:
    return _n(MARK.sub("", s))


def _body_texts(doc) -> list[str]:
    """Тексты абзацев верхнего уровня до раздела примечаний."""
    out = []
    for p in doc.paragraphs:
        if p.text.strip().startswith(NOTES_HEAD):
            break
        out.append(p.text)
    return out


def _has_notes(doc) -> bool:
    return any(p.text.strip().startswith(NOTES_HEAD) for p in doc.paragraphs)


def _grif_tables(doc) -> list:
    return [t._tbl for t in doc.tables if "УТВЕРЖДАЮ" in "".join(t._tbl.itertext())]


def _table_pars(doc) -> list[str]:
    """Абзацы внутри таблиц, кроме таблицы с грифом утверждения (её отдельное замечание)."""
    skip = _grif_tables(doc)
    res = []
    for p in doc.element.body.iter(qn("w:p")):
        if p.getparent().tag == qn("w:body"):
            continue
        if any(t in p.iterancestors() for t in skip):
            continue
        res.append(Paragraph(p, doc._body).text)
    return res


# ---------------------------------------------------------------- шаблоны
def _to_regex(t: str) -> re.Pattern:
    if t in FREE_PARAGRAPHS:
        return re.compile(r".+")
    tok = re.compile(
        r"\[[^\]]*\]"                       # [необязательный фрагмент] и [проекту / проектам]
        r"|_{3,}|И\. О\. Фамилия"
        r"|([а-яё]+)\(([а-яё]+)\)"          # реализуемого(ых)
        r"|([а-яё]+) / ([а-яё]+)"           # вариант без скобок
    )
    parts, pos = [], 0
    for m in tok.finditer(t):
        parts.append(re.escape(t[pos:m.start()]))
        s = m.group(0)
        if s.startswith("["):
            parts.append(r".*?")
        elif m.group(1):
            parts.append(r"[а-яё]+")
        elif m.group(3):
            parts.append(f"(?:{re.escape(m.group(3))}|{re.escape(m.group(4))})")
        else:
            parts.append(r".+?")
        pos = m.end()
    parts.append(re.escape(t[pos:]))
    return re.compile("".join(parts), re.S)


def _patterns(path: Path) -> list[tuple[str, re.Pattern]]:
    doc = Document(str(path))
    out = []
    for text in _body_texts(doc):
        t = _clean(text)
        if not t or t == "МОДУЛЬНЫЙ ДОКУМЕНТ":
            continue
        out.append((t, _to_regex(t)))
    return out


def _doc_for_match(text: str) -> str:
    return _clean(VARIANT.sub(r"\1", text))


def _match(doc_pars: list[str], patterns) -> tuple[list[str], list[str]]:
    """Возвращает (абзацы документа без пары в модуле, абзацы модуля без пары)."""
    matched: set[int] = set()
    extra = []
    pos = 0
    for raw in doc_pars:
        d = _doc_for_match(raw)
        if not d or d == "МОДУЛЬНЫЙ ДОКУМЕНТ":
            continue
        pos += 1
        hits = [i for i, (t, rx) in enumerate(patterns) if t not in FREE_PARAGRAPHS and rx.fullmatch(d)]
        if not hits and pos <= 8 and len(d) <= 100 and not d.endswith("."):
            # «свободные» строки шапки (наименование организации, город) подставляются своими словами
            hits = [i for i, (t, _) in enumerate(patterns) if t in FREE_PARAGRAPHS and i not in matched]
        if not hits:
            extra.append(d)
            continue
        matched.add(next((i for i in hits if i not in matched), hits[0]))
    missing = [patterns[i][0] for i in range(len(patterns)) if i not in matched]
    return extra, missing


def detect_module(doc, templates: dict[str, Path]) -> tuple[str | None, float]:
    pars = _body_texts(doc)
    best, best_score = None, 0.0
    for mid, path in templates.items():
        pats = _patterns(path)
        extra, missing = _match(pars, pats)
        n_doc = max(1, len([p for p in pars if _clean(p)]))
        score = 0.5 * (1 - len(extra) / n_doc) + 0.5 * (1 - len(missing) / max(1, len(pats)))
        if score > best_score:
            best, best_score = mid, score
    return (best, round(best_score, 2)) if best_score >= 0.5 else (None, round(best_score, 2))


# ---------------------------------------------------------------- проверки
def _issue(sev: str, title: str, detail: str = "", where: str = "", refs: list[str] | None = None) -> dict:
    return {"severity": sev, "title": title, "detail": detail, "where": where, "refs": refs or []}


def _snippet(s: str, n: int = 110) -> str:
    s = _n(MARK.sub("", s))
    return s if len(s) <= n else s[: n - 1] + "…"


def _text_date(m: re.Match) -> date | None:
    try:
        return date(int(m[3]), MONTHS.index(m[2]) + 1, int(m[1]))
    except (ValueError, IndexError):
        return None


def _check_format(doc) -> list[dict]:
    out = []
    s = doc.sections[0]
    mm = lambda v: round(v.mm, 1) if v is not None else None  # noqa: E731
    left, right, top, bottom = (mm(s.left_margin), mm(s.right_margin), mm(s.top_margin), mm(s.bottom_margin))
    bad = []
    if left is not None and abs(left - 30) > 0.6:
        bad.append(f"левое {left} мм (нужно 30)")
    if right is not None and right < 19.4:
        bad.append(f"правое {right} мм (не менее 20)")
    if top is not None and top < 24.4:
        bad.append(f"верхнее {top} мм (не менее 25)")
    if bottom is not None and bottom < 14.4:
        bad.append(f"нижнее {bottom} мм (не менее 15)")
    if bad:
        out.append(_issue("warn", "Поля страницы не соответствуют Инструкции", "; ".join(bad), refs=[f"{INS}, п. 8.16"]))
    fonts = {r.font.name for p in doc.paragraphs for r in p.runs if r.font.name and r.text.strip()}
    if fonts - {"Times New Roman"}:
        out.append(_issue("warn", "Основной шрифт должен быть Times New Roman",
                          "Найдено: " + ", ".join(sorted(fonts - {"Times New Roman"})), refs=[f"{INS}, п. 8.16"]))
    return out


IS_ELEMENTS = [  # п. 9.5.6 Инструкции: (что указывает приказ, шаблон поиска)
    ("режим эксплуатации", r"режим"),
    ("условия доступа", r"доступ"),
    ("порядок миграции данных", r"миграц"),
    ("порядок передачи документации", r"передач\w*\s+(эксплуатационн\w+\s+)?документац"),
    ("действия по информационной безопасности", r"безопасност"),
]


def check_docx(data: bytes, templates: dict[str, Path], module_id: str | None = None) -> dict:
    doc = Document(io.BytesIO(data))
    body = _body_texts(doc)
    tbl = _table_pars(doc)
    all_text = "\n".join(body + tbl)
    issues: list[dict] = []

    if not module_id or module_id not in templates:
        module_id, score = detect_module(doc, templates)
    else:
        score = 1.0

    # --- служебные элементы модуля, которые должны быть удалены (МР 3.3)
    if _grif_tables(doc) or "УТВЕРЖДАЮ" in all_text:
        issues.append(_issue("error", "Остался гриф утверждения «УТВЕРЖДАЮ» вверху первой страницы",
                             "В проекте приказа его нужно удалить.", refs=[f"{MR}, п. 3.3"]))
    if any(_n(t) == "МОДУЛЬНЫЙ ДОКУМЕНТ" for t in body):
        issues.append(_issue("error", "Осталась пометка «МОДУЛЬНЫЙ ДОКУМЕНТ»",
                             "Её нужно удалить вместе с грифом утверждения.", refs=[f"{MR}, п. 3.3"]))
    if _has_notes(doc):
        issues.append(_issue("error", "Остался раздел «Примечания по заполнению»",
                             "Пояснения к модулю в подписываемом документе не оставляют.", refs=[f"{MR}, п. 3.3"]))

    # --- незаполненные поля
    for t in body + [x for x in tbl if x.strip("_ \t")]:
        if BLANK.search(t):
            issues.append(_issue("error", "Не заполнено поле (осталась линия «____»)", where=_snippet(t),
                                 refs=[f"{MR}, п. 3.3", f"{MR}, п. 3.2"]))
    for t in body + tbl:
        if VARIANT.search(t):
            issues.append(_issue("error", "Не выбран вариант (остался знак «/»)", where=_snippet(t),
                                 detail="Неиспользуемый вариант и знак «/» нужно удалить.", refs=[f"{MR}, п. 3.3"]))
    marks = sorted({m for t in body + tbl for m in re.findall(r"\[(\d{1,2})\]", t)}, key=int)
    if marks:
        issues.append(_issue("error", "Остались номера примечаний: " + ", ".join(f"[{m}]" for m in marks),
                             "Номера примечаний нужно удалить.", refs=[f"{MR}, п. 3.3"]))
    if re.search(r"И\. О\. Фамилия", all_text):
        issues.append(_issue("warn", "Не указан подписант (осталось «И. О. Фамилия»)",
                             "Подпись включает должность, инициалы и фамилию.", refs=[f"{INS}, п. 8.19.2"]))

    # --- обязательные реквизиты
    kind_word = "РАСПОРЯЖЕНИЕ" if (module_id or "").startswith("decree_") else "ПРИКАЗ"
    verb = "ОБЯЗЫВАЮ:" if kind_word == "РАСПОРЯЖЕНИЕ" else "ПРИКАЗЫВАЮ:"
    verb_ref = f"{INS}, п. " + ("9.7.4" if kind_word == "РАСПОРЯЖЕНИЕ" else "9.5.4")
    norm_body = [_n(t) for t in body]
    if kind_word not in norm_body:
        issues.append(_issue("error", f"Нет наименования вида документа «{kind_word}»", refs=[f"{INS}, п. 8.6"]))
    if verb not in norm_body:
        issues.append(_issue("error", f"Нет слова «{verb}» отдельной строкой прописными буквами",
                             "Им начинается распорядительная часть.", refs=[verb_ref]))
    if not any("РЕГИСТРАЦИОННЫЙ НОМЕР" in t for t in body):
        issues.append(_issue("error", "Удалена зона даты и регистрационного номера",
                             "Служебные зоны с пометкой «НЕ УДАЛЯТЬ» должны сохраняться.",
                             refs=[f"{MR}, п. 3.3", f"{INS}, п. 8.7–8.8"]))
    if "ВИЗУАЛИЗАЦИЯ ЭП" not in all_text:
        issues.append(_issue("error", "Удалён блок визуализации электронной подписи",
                             "Блок с пометкой «НЕ УДАЛЯТЬ» должен сохраняться.", refs=[f"{MR}, п. 3.3", f"{INS}, п. 8.20"]))
    if not any(re.match(r"^(О|Об)\s+\S+\s+\S+", _n(t)) and not BLANK.search(t) for t in body):
        issues.append(_issue("error", "Нет заголовка документа («О …» / «Об …»)", refs=[f"{INS}, п. 8.15"]))

    # --- даты
    seen = set()
    for d_, m_, y_ in DATE_IN_TEXT.findall(all_text):
        if (d_, m_, y_) not in seen and not _parse_date(f"{d_}.{m_}.{y_}"):
            seen.add((d_, m_, y_))
            issues.append(_issue("error", f"Несуществующая дата {d_}.{m_}.{y_}", refs=[f"{INS}, п. 8.7"]))
    for m in TEXT_DATE.finditer(all_text):
        if m[2] in MONTHS and not _text_date(m):
            issues.append(_issue("error", f"Несуществующая дата «{m[1]}» {m[2]} {m[3]} г.", refs=[f"{INS}, п. 8.7"]))

    if module_id == "order_is_commissioning":
        base = next((t for t in body if t.startswith("На основании")), "")
        intro = next((t for t in body if t.startswith("1. Ввести")), "")
        rd = [d for d in (_parse_date(".".join(m)) for m in DATE_IN_TEXT.findall(base)) if d]
        md = TEXT_DATE.search(intro)
        sd = _text_date(md) if md and md[2] in MONTHS else None
        if rd and sd and sd < max(rd):
            issues.append(_issue("error", "Дата ввода в эксплуатацию раньше даты документа о готовности",
                                 f"Ввод с {sd:%d.%m.%Y}, документ о готовности от {max(rd):%d.%m.%Y}.",
                                 refs=["Модуль, примечание 8"]))
        if base and not BLANK.search(base) and not rd and not DATE_IN_TEXT.search(base):
            issues.append(_issue("warn", "В основании приказа не видно даты документа о готовности",
                                 refs=["Модуль, примечание 2", f"{INS}, п. 9.5.6"]))

    # --- единообразие сокращённого наименования
    m_short = re.search(r"далее\s*-\s*«([^»]+)»", all_text)
    if m_short and "_" not in m_short[1]:
        short = m_short[1].strip()
        pre = QUOTED.findall(all_text[: m_short.start()])
        full = pre[-1].strip() if pre else ""
        odd = []
        for q in QUOTED.findall(all_text[m_short.end():]):
            q = q.strip()
            if q and "_" not in q and not q.isdigit() and q not in (short, full) and q not in odd:
                odd.append(q)
        for q in odd:
            issues.append(_issue("warn", f"Наименование «{q}» не совпадает с сокращённым «{short}»",
                                 "Сокращённое наименование должно быть одинаковым по всему приказу.",
                                 refs=[f"{MR}, п. 3.3", "Модуль, примечания 4, 7, 10, 12, 13"]))

    # --- сверка с модулем
    differs = False
    if module_id:
        extra, missing = _match(body, _patterns(templates[module_id]))
        flagged = lambda t: bool(BLANK.search(t) or "/" in t)  # noqa: E731
        extra = [e for e in extra if not flagged(e) and e != "МОДУЛЬНЫЙ ДОКУМЕНТ"]
        for e in extra:
            differs = True
            issues.append(_issue("warn", "Текст отличается от модуля или добавлен", where=_snippet(e),
                                 detail="Постоянные формулировки модуля менять нельзя без полного согласования.",
                                 refs=[f"{MR}, п. 1.3"]))
        for mi in missing:
            if BLANK.search(mi) or "/" in mi:
                continue
            differs = True
            issues.append(_issue("warn", "Нет пункта модуля", where=_snippet(mi), refs=[f"{MR}, п. 1.3"]))

        if module_id == "order_is_commissioning":
            lower = all_text.lower()
            absent = [name for name, rx in IS_ELEMENTS if not re.search(rx, lower)]
            if absent:
                issues.append(_issue(
                    "info", "Не найдены сведения, которые обычно указывают в приказе о вводе ИС",
                    "Нет упоминания: " + "; ".join(absent) + ". Модуль их не требует, но предлагает дополнять приказ при необходимости.",
                    refs=[f"{INS}, п. 9.5.6", "Модуль, раздел «Общие правила»"]))

    issues.extend(_check_format(doc))

    order = {"error": 0, "warn": 1, "info": 2}
    issues.sort(key=lambda i: order[i["severity"]])
    counts = {k: sum(1 for i in issues if i["severity"] == k) for k in order}
    return {
        "module_id": module_id,
        "module_confidence": score,
        "differs_from_module": differs,
        "counts": counts,
        "issues": issues,
    }


# ---------------------------------------------------------------- дальнейшие шаги
def next_steps(result: dict, visa: dict | None) -> list[dict]:
    mid = result["module_id"] or ""
    decree = mid.startswith("decree_")
    steps: list[dict] = []
    if result["counts"]["error"]:
        steps.append({"title": "Исправьте ошибки и загрузите проект ещё раз",
                      "detail": "Пока в документе есть ошибки, его нельзя направлять по маршруту.",
                      "refs": [f"{MR}, п. 3.3"]})
    if mid:
        if result["differs_from_module"]:
            steps.append({"title": "Проект отличается от модуля: нужно полное согласование",
                          "detail": "Если изменены постоянные формулировки, обязательные пункты или смысл решения, проект проходит полное согласование по главе 10 Инструкции.",
                          "refs": [f"{MR}, п. 1.3", f"{INS}, гл. 10"]})
        elif not result["counts"]["error"]:
            steps.append({"title": "Проект соответствует модулю",
                          "detail": "Подразделениям, согласовавшим модуль, повторное полное согласование не направляется. Лингвистическая экспертиза и визирование Департаментом ДОУ сохраняются.",
                          "refs": [f"{MR}, п. 1.3"]})
    steps.append({"title": "Создайте регистрационную карточку проекта в ИИС ДО",
                  "detail": "В «Примечании» укажите, что проект подготовлен на основе модуля (вид, наименование, номер, дата, редакция), и свяжите карточки. Регистрационный номер до подписания не присваивается.",
                  "refs": [f"{MR}, п. 2.1–2.4"]})
    items = []
    if visa:
        items = list(visa.get("required_visas", [])) + [
            x["dept"] + (f" (если: {x['condition']})" if x.get("condition") else "") for x in visa.get("additional_visas", [])]
    steps.append({"title": "Согласование и визирование",
                  "detail": "Визируйте строго по порядку: ответственный исполнитель, участники разработки, руководитель подразделения, заинтересованные подразделения, подразделения обязательного визирования по Матрице, юридическое подразделение, Департамент ДОУ, координирующее должностное лицо."
                            + (" Обязательные визы по Матрице для этого документа: " + "; ".join(items) + "." if items else ""),
                  "refs": [f"{INS}, п. 10.2", f"{INS}, п. 10.6", f"{INS}, п. 10.11"]})
    steps.append({"title": "Срок рассмотрения",
                  "detail": "Максимальный срок каждого первичного и повторного рассмотрения: "
                            + ("3 рабочих дня (распоряжение)." if decree else "10 рабочих дней (приказ по основной деятельности).")
                            + " Это срок на каждый круг рассмотрения, а не на всё согласование.",
                  "refs": [f"{INS}, п. 10.3"]})
    steps.append({"title": "Лингвистическая экспертиза",
                  "detail": "Приказы по основной деятельности и распоряжения проходят экспертизу в Департаменте ДОУ (русский язык, терминология, стиль, оформление).",
                  "refs": [f"{INS}, п. 10.13", f"{MR}, п. 4.2"]})
    steps.append({"title": "Пояснительная записка",
                  "detail": "Нужна, если объём больше трёх страниц или документ вносит изменения. Оформляется отдельным файлом.",
                  "refs": [f"{INS}, п. 10.12"]})
    steps.append({"title": "Подписание",
                  "detail": "Обычно не более трёх рабочих дней после направления подписанту. Завизированная и подписываемая версии должны совпадать, после подписания исправления запрещены.",
                  "refs": [f"{INS}, п. 10.15", f"{INS}, п. 8.19", f"{MR}, п. 5.1"]})
    steps.append({"title": "Регистрация и рассылка",
                  "detail": "Регистрация не позднее рабочего дня, следующего за днём подписания. Перед присвоением номера Департамент ДОУ проверяет подпись, маршрут и визы, реквизиты и файл. Список рассылки формирует ответственный исполнитель до регистрации.",
                  "refs": [f"{INS}, п. 7.1", f"{INS}, п. 7.3", f"{INS}, п. 7.14", f"{INS}, п. " + ("9.8" if decree else "9.5.10")]})
    return steps